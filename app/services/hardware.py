import csv
import io
import logging
from datetime import datetime, timezone
from typing import Optional, List, Dict, Any, Tuple
import pandas as pd
import qrcode
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy.orm import Session

from app.models.hardware import Hardware, StatusEnum, ModelEnum
from app.core.config import settings

logger = logging.getLogger(__name__)


class HardwareService:
    def __init__(self, db: Session):
        self.db = db
    
    def get_hardware_by_id(self, hardware_id: int) -> Optional[Hardware]:
        return self.db.query(Hardware).filter(Hardware.id == hardware_id).first()
    
    def get_hardware_by_serial(self, serial_number: str) -> Optional[Hardware]:
        if not serial_number:
            return None
        return (
            self.db.query(Hardware)
            .filter(Hardware.serial_number == serial_number)
            .first()
        )
    
    def get_filtered_hardware_query(self, 
                                   search: Optional[str] = None,
                                   status: Optional[List[str]] = None,
                                   model: Optional[str] = None,
                                   center: Optional[str] = None):
        query = self.db.query(Hardware)

        if search:
            search_term = f"%{search}%"
            query = query.filter(
                (Hardware.hostname.ilike(search_term))
                | (Hardware.ip.ilike(search_term))
                | (Hardware.mac.ilike(search_term))
                | (Hardware.serial_number.ilike(search_term))
                | (Hardware.uuid.ilike(search_term))
                | (Hardware.enduser.ilike(search_term))
                | (Hardware.ticket.ilike(search_term))
                | (Hardware.po_ticket.ilike(search_term))
                | (Hardware.center.ilike(search_term))
                | (Hardware.comment.ilike(search_term))
                | (Hardware.admin.ilike(search_term))
            )

        status_filter_list = []
        if status and len(status) > 0:
            for s in status:
                if s:
                    try:
                        status_enum = StatusEnum(s)
                        status_filter_list.append(status_enum)
                    except ValueError:
                        pass
        
        if status_filter_list:
            query = query.filter(Hardware.status.in_(status_filter_list))
        else:
            default_statuses = [s for s in StatusEnum if s != StatusEnum.COMPLETED]
            query = query.filter(Hardware.status.in_(default_statuses))
        
        if model:
            try:
                model_enum = ModelEnum(model)
                query = query.filter(Hardware.model == model_enum)
            except ValueError:
                pass
        
        if center:
            query = query.filter(Hardware.center.ilike(f"%{center}%"))
        
        return query, status_filter_list if status_filter_list else [s for s in StatusEnum if s != StatusEnum.COMPLETED]
    
    def get_hardware_list(self,
                         search: Optional[str] = None,
                         status: Optional[List[str]] = None,
                         model: Optional[str] = None,
                         center: Optional[str] = None,
                         page: int = 1,
                         per_page: int = 20,
                         sort_by: str = "updated_at",
                         sort_order: str = "desc") -> Dict[str, Any]:

        query, status_filter_list = self.get_filtered_hardware_query(search, status, model, center)
        
        total_count = query.count()
        
        sort_column = getattr(Hardware, sort_by, Hardware.updated_at)
        if sort_order.lower() == "asc":
            query = query.order_by(sort_column.asc())
        else:
            query = query.order_by(sort_column.desc())
        
        offset = (page - 1) * per_page
        hardware_list = query.offset(offset).limit(per_page).all()
        
        total_pages = (total_count + per_page - 1) // per_page
        
        status_counts = {}
        for s in StatusEnum:
            count = self.db.query(Hardware).filter(Hardware.status == s).count()
            status_counts[s.value] = count
        
        return {
            "hardware_list": hardware_list,
            "total_count": total_count,
            "current_page": page,
            "per_page": per_page,
            "total_pages": total_pages,
            "status_counts": status_counts,
            "status_filter": [s.value for s in status_filter_list]
        }
    
    def create_hardware(self, hardware_data: Dict[str, Any], current_user: Dict[str, Any]) -> Hardware:
        now = datetime.now(timezone.utc)
        
        hardware = Hardware(
            hostname=hardware_data.get('hostname'),
            serial_number=hardware_data.get('serial_number'),
            model=hardware_data.get('model'),
            status=hardware_data.get('status'),
            ip=hardware_data.get('ip'),
            mac=hardware_data.get('mac'),
            uuid=hardware_data.get('uuid'),
            center=hardware_data.get('center'),
            enduser=hardware_data.get('enduser'),
            ticket=hardware_data.get('ticket'),
            po_ticket=hardware_data.get('po_ticket'),
            admin=current_user["username"],
            comment=hardware_data.get('comment'),
            missing=hardware_data.get('missing', False),
            created_at=now,
            updated_at=now,
            shipped_at=now if hardware_data.get('status') == StatusEnum.SHIPPED else None,
        )
        
        self.db.add(hardware)
        self.db.commit()
        self.db.refresh(hardware)
        
        return hardware
    
    def update_hardware(self, hardware_id: int, hardware_data: Dict[str, Any], current_user: Dict[str, Any]) -> Hardware:
        hardware = self.get_hardware_by_id(hardware_id)
        if not hardware:
            raise ValueError("Hardware not found")
        
        old_status = hardware.status
        
        hardware.hostname = hardware_data.get('hostname')
        hardware.serial_number = hardware_data.get('serial_number')
        hardware.model = hardware_data.get('model')
        hardware.status = hardware_data.get('status')
        hardware.ip = hardware_data.get('ip')
        hardware.mac = hardware_data.get('mac')
        hardware.uuid = hardware_data.get('uuid')
        hardware.center = hardware_data.get('center')
        hardware.enduser = hardware_data.get('enduser')
        hardware.ticket = hardware_data.get('ticket')
        hardware.po_ticket = hardware_data.get('po_ticket')
        hardware.admin = current_user["username"]
        hardware.comment = hardware_data.get('comment')
        hardware.missing = hardware_data.get('missing', False)
        hardware.updated_at = datetime.now(timezone.utc)
        
        if hardware_data.get('status') == StatusEnum.SHIPPED and old_status != StatusEnum.SHIPPED:
            hardware.shipped_at = datetime.now(timezone.utc)
        
        self.db.commit()
        return hardware
    
    def delete_hardware(self, hardware_id: int) -> bool:
        hardware = self.get_hardware_by_id(hardware_id)
        if not hardware:
            raise ValueError("Hardware not found")
        
        self.db.delete(hardware)
        self.db.commit()
        return True
    
    def export_hardware_to_excel(self,
                                 search: Optional[str] = None,
                                 status: Optional[List[str]] = None,
                                 model: Optional[str] = None,
                                 center: Optional[str] = None) -> io.BytesIO:

        query, _ = self.get_filtered_hardware_query(search, status, model, center)
        hardware_list = query.order_by(Hardware.updated_at.desc()).all()

        data = []
        for hw in hardware_list:
            data.append({
                'ID': hw.id,
                'Hostname': hw.hostname,
                'Serial Number': hw.serial_number,
                'Model': hw.model.value if hw.model else '',
                'Status': hw.status.value if hw.status else '',
                'IP Address': hw.ip or '',
                'MAC Address': hw.mac or '',
                'UUID': hw.uuid or '',
                'Center': hw.center or '',
                'End User': hw.enduser or '',
                'Ticket': hw.ticket or '',
                'PO Ticket': hw.po_ticket or '',
                'Admin': hw.admin,
                'Comment': hw.comment or '',
                'Missing': 'Yes' if hw.missing else 'No',
                'Created At': hw.created_at.strftime('%Y-%m-%d %H:%M:%S') if hw.created_at else '',
                'Updated At': hw.updated_at.strftime('%Y-%m-%d %H:%M:%S') if hw.updated_at else '',
                'Shipped At': hw.shipped_at.strftime('%Y-%m-%d %H:%M:%S') if hw.shipped_at else '',
            })

        if not data:
            return io.BytesIO()

        df = pd.DataFrame(data)
        
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Hardware Inventory', index=False)

            worksheet = writer.sheets['Hardware Inventory']

            num_rows, num_cols = df.shape

            last_col_letter = get_column_letter(num_cols)
            table_ref = f"A1:{last_col_letter}{num_rows + 1}"

            display_name = f"HardwareTabelle_{datetime.now().strftime('%Y%m%d%H%M%S')}"
            
            table = Table(displayName=display_name, ref=table_ref)

            style = TableStyleInfo(
                name="TableStyleMedium9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            table.tableStyleInfo = style

            worksheet.add_table(table)
            
            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter
                for cell in column_cells:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                
                adjusted_width = min(max_length + 2, 50)
                worksheet.column_dimensions[column_letter].width = adjusted_width
        
        output.seek(0)
        return output
    
    def import_hardware_from_file(self, file_content: bytes, filename: str, current_user: Dict[str, Any]) -> Dict[str, Any]:
        parsed = self.parse_import_file(file_content, filename)

        if not parsed["valid_items"]:
            return {
                "total_rows": parsed["total_rows"],
                "created": 0,
                "updated": 0,
                "errors": parsed["errors"],
            }

        apply_summary = self.apply_import_data(
            [item["data"] for item in parsed["valid_items"]],
            current_user,
        )

        return {
            "total_rows": parsed["total_rows"],
            "created": apply_summary["created"],
            "updated": apply_summary["updated"],
            "errors": parsed["errors"] + apply_summary["errors"],
        }

    def parse_import_file(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        filename = filename.lower()

        if filename.endswith('.csv'):
            csv_string = file_content.decode('utf-8')
            csv_reader = csv.DictReader(io.StringIO(csv_string))
            rows = list(csv_reader)
        elif filename.endswith(('.xlsx', '.xls')):
            df = pd.read_excel(io.BytesIO(file_content))
            rows = df.to_dict('records')
        else:
            raise ValueError("File must be a CSV or Excel file (.csv, .xlsx, .xls)")

        total_rows = len(rows)
        errors: List[str] = []
        valid_items: List[Dict[str, Any]] = []
        seen_serials: set[str] = set()

        for row_num, row in enumerate(rows, start=2):
            try:
                def safe_get(key, default=''):
                    value = row.get(key, default)
                    if pd.isna(value):
                        return default
                    return str(value).strip()

                hardware_data = self._extract_hardware_from_row(row, safe_get)

                if not all([
                    hardware_data.get('hostname'),
                    hardware_data.get('serial_number'),
                    hardware_data.get('model'),
                    hardware_data.get('status'),
                ]):
                    errors.append(f"Row {row_num}: Missing required fields")
                    continue

                serial = hardware_data['serial_number']
                if serial in seen_serials:
                    errors.append(f"Row {row_num}: Duplicate serial number '{serial}' in file")
                    continue
                seen_serials.add(serial)

                try:
                    model_enum = ModelEnum(hardware_data['model'])
                except ValueError:
                    errors.append(f"Row {row_num}: Invalid model '{hardware_data['model']}'")
                    continue

                try:
                    status_enum = StatusEnum(hardware_data['status'])
                except ValueError:
                    errors.append(f"Row {row_num}: Invalid status '{hardware_data['status']}'")
                    continue

                hardware_data['model'] = model_enum.value
                hardware_data['status'] = status_enum.value

                existing = self.get_hardware_by_serial(serial)
                action = 'update' if existing else 'create'

                valid_items.append({
                    'row': row_num,
                    'serial_number': serial,
                    'hostname': hardware_data.get('hostname'),
                    'model': hardware_data.get('model'),
                    'status': hardware_data.get('status'),
                    'action': action,
                    'data': hardware_data,
                })

            except Exception as exc:
                errors.append(f"Row {row_num}: {exc}")

        create_count = sum(1 for item in valid_items if item['action'] == 'create')
        update_count = sum(1 for item in valid_items if item['action'] == 'update')

        return {
            'total_rows': total_rows,
            'valid_items': valid_items,
            'errors': errors,
            'create_count': create_count,
            'update_count': update_count,
        }

    def apply_import_data(self, items: List[Dict[str, Any]], current_user: Dict[str, Any]) -> Dict[str, Any]:
        created = 0
        updated = 0
        errors: List[str] = []

        for item in items:
            try:
                hardware_data = item.copy()
                hardware_data['model'] = ModelEnum(hardware_data['model'])
                hardware_data['status'] = StatusEnum(hardware_data['status'])

                action, _ = self.upsert_hardware_by_serial(hardware_data, current_user)
                if action == 'created':
                    created += 1
                else:
                    updated += 1
            except Exception as exc:
                serial = item.get('serial_number', 'UNKNOWN')
                errors.append(f"Serial {serial}: {exc}")

        return {'created': created, 'updated': updated, 'errors': errors}

    def upsert_hardware_by_serial(
        self, hardware_data: Dict[str, Any], current_user: Dict[str, Any]
    ) -> Tuple[str, Hardware]:
        serial_number = hardware_data.get('serial_number')
        if not serial_number:
            raise ValueError('serial_number is required')

        existing = self.get_hardware_by_serial(serial_number)
        if existing:
            updated = self.update_hardware(existing.id, hardware_data, current_user)
            return 'updated', updated

        created = self.create_hardware(hardware_data, current_user)
        return 'created', created
    
    def _extract_hardware_from_row(self, row: Dict[str, Any], safe_get) -> Dict[str, Any]:
        missing_value = safe_get('Missing').lower()
        missing = missing_value in ['yes', 'true', '1']
        
        return {
            'hostname': safe_get('Hostname'),
            'serial_number': safe_get('Serial Number'),
            'model': safe_get('Model'),
            'status': safe_get('Status'),
            'ip': safe_get('IP Address') or None,
            'mac': safe_get('MAC Address') or None,
            'uuid': safe_get('UUID') or None,
            'center': safe_get('Center') or None,
            'enduser': safe_get('End User') or None,
            'ticket': safe_get('Ticket') or None,
            'po_ticket': safe_get('PO Ticket') or None,
            'comment': safe_get('Comment') or None,
            'missing': missing,
        }
    
    def change_hardware_status(self, hardware_id: int, status: StatusEnum, current_user: Dict[str, Any]) -> Hardware:
        hardware = self.get_hardware_by_id(hardware_id)
        if not hardware:
            raise ValueError("Hardware not found")
        
        old_status = hardware.status
        hardware.status = status
        hardware.admin = current_user["username"]
        hardware.updated_at = datetime.now(timezone.utc)
        
        if status == StatusEnum.SHIPPED and old_status != StatusEnum.SHIPPED:
            hardware.shipped_at = datetime.now(timezone.utc)
        
        self.db.commit()
        return hardware
    
    def cycle_hardware_status(self, hardware_id: int, current_user: Dict[str, Any]) -> Tuple[Hardware, str]:
        hardware = self.get_hardware_by_id(hardware_id)
        if not hardware:
            raise ValueError("Hardware not found")
        
        status_cycle = [
            StatusEnum.IN_STOCK,
            StatusEnum.RESERVED,
            StatusEnum.IMAGING,
            StatusEnum.SHIPPED,
            StatusEnum.COMPLETED
        ]
        
        try:
            current_index = status_cycle.index(hardware.status)
            if current_index == len(status_cycle) - 1:
                raise ValueError("Device is already completed and cannot be cycled further")
            
            new_status = status_cycle[current_index + 1]
        except ValueError:
            new_status = status_cycle[0]
        
        hardware = self.change_hardware_status(hardware_id, new_status, current_user)
        
        status_display = {
            'IN_STOCK': 'In Stock',
            'RESERVED': 'Reserved', 
            'IMAGING': 'Imaging',
            'SHIPPED': 'Shipped',
            'COMPLETED': 'Completed'
        }.get(new_status.value, new_status.value)
        
        return hardware, status_display
    
    def generate_qr_code(self, hardware_id: int) -> Tuple[io.BytesIO, str]:
        hardware = self.get_hardware_by_id(hardware_id)
        if not hardware:
            raise ValueError("Hardware not found")
        
        api_url = f"{settings.base_url}/hardware/{hardware_id}"
        
        qr = qrcode.QRCode(
            version=1,
            error_correction=qrcode.constants.ERROR_CORRECT_L,
            box_size=10,
            border=4,
        )
        qr.add_data(api_url)
        qr.make(fit=True)
        
        img = qr.make_image(fill_color="black", back_color="white")
        
        img_buffer = io.BytesIO()
        img.save(img_buffer, format='PNG')
        img_buffer.seek(0)
        
        serial = hardware.serial_number.replace('/', '-').replace('\\', '-')
        filename = f"QR_{serial}_{hardware.hostname}.png"
        
        return img_buffer, filename

    def generate_label_csv(self, hardware_id: int) -> Tuple[str, str]:
        hardware = self.get_hardware_by_id(hardware_id)
        if not hardware:
            raise ValueError("Hardware not found")

        csv_content = []
        csv_content.append(f"HN,{hardware.hostname or ''}")
        csv_content.append(f"SN,{hardware.serial_number or ''}")
        csv_content.append(f"IP,{hardware.ip or ''}")
        csv_content.append(f"T#,{hardware.ticket or ''}")
        csv_content.append(f"PO#,{hardware.po_ticket or ''}")
        csv_content.append(f"User,{hardware.enduser or ''}")
        csv_content.append(f"Cent,{hardware.center or ''}")
        
        csv_string = '\n'.join(csv_content)
        
        serial = hardware.serial_number.replace('/', '-').replace('\\', '-')
        filename = f"label_{serial}_{hardware.hostname}.csv"
        
        return csv_string, filename
